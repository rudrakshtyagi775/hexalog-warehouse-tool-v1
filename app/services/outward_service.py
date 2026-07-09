import csv
import io
from datetime import UTC, date, datetime, timedelta

import structlog
from fastapi import HTTPException, status
from sqlalchemy import func as sql_func
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.enums import (
    AuditModuleEnum,
    CounterTypeEnum,
    LedgerSourceTypeEnum,
    OutwardBoxStatusEnum,
    OutwardPoStatusEnum,
    OutwardScanResultEnum,
)
from app.models.inward import InventoryLedgerEntry
from app.models.outward import OutwardBox, OutwardPO, OutwardPOLine, OutwardScan
from app.schemas.outward import (
    LabelGenerateResponse,
    LabelHistoryItem,
    LabelHistoryResponse,
    OpenPOListResponse,
    OpenPOSummary,
    OutwardBoxListResponse,
    OutwardBoxSummary,
    OutwardPOPreviewResponse,
    OutwardPOPreviewRow,
)
from app.services.audit_service import write_audit_log
from app.services.inward_service import _next_counter

log = structlog.get_logger(__name__)

_PO_REQUIRED_COLUMNS = {"po_number", "ean", "ordered_qty"}
_MAX_ALLOC_RETRIES = 5


def _read_upload_rows(filename: str, file_bytes: bytes) -> tuple[list[str] | None, list[dict]]:
    """Parse an uploaded PO file (CSV or XLSX, per OUT-1) into (fieldnames, rows),
    mirroring csv.DictReader's shape so downstream validation stays format-agnostic.

    fieldnames is None when the file has no header row (mirrors DictReader).
    """
    if filename.lower().endswith((".xlsx", ".xls")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if header_row is None or all(cell is None for cell in header_row):
            return None, []

        fieldnames = [str(cell).strip() if cell is not None else "" for cell in header_row]
        rows: list[dict] = []
        for raw_row in rows_iter:
            if all(cell is None for cell in raw_row):
                continue
            row: dict = {}
            for key, value in zip(fieldnames, raw_row):
                if value is None:
                    row[key] = ""
                elif isinstance(value, float) and value.is_integer():
                    row[key] = str(int(value))  # avoid "10.0" for whole-number cells
                else:
                    row[key] = str(value)
            rows.append(row)
        return fieldnames, rows

    text_io = io.StringIO(file_bytes.decode("utf-8-sig"))
    reader = csv.DictReader(text_io)
    return reader.fieldnames, list(reader)


def _consolidate_lines(lines_data: list[dict]) -> tuple[list[dict], int]:
    """Sum ordered_qty for duplicate EANs into a single line (OUT-3).

    OutwardPOLine has a unique (outward_po_id, ean) constraint, so an
    unconsolidated CSV with a repeated EAN would otherwise fail at commit.
    Returns (consolidated_lines, duplicate_row_count).
    """
    merged: dict[str, dict] = {}
    duplicate_rows = 0
    for ld in lines_data:
        existing = merged.get(ld["ean"])
        if existing is None:
            merged[ld["ean"]] = dict(ld)
        else:
            existing["ordered_qty"] += ld["ordered_qty"]
            existing["description"] = existing["description"] or ld["description"]
            duplicate_rows += 1
    return list(merged.values()), duplicate_rows


# ── PO upload ─────────────────────────────────────────────────────────────────

async def upload_po(
    db: AsyncSession,
    *,
    customer_id: int,
    organisation_id: int,
    uploaded_by: int,
    ip_address: str | None,
    po_number: str,
    filename: str = "po.csv",
    csv_bytes: bytes,
) -> OutwardPO:
    """Parse CSV/XLSX (OUT-1), validate columns, detect duplicates, create OutwardPO + lines."""
    from app.models.customer import Customer

    # 1. Parse the uploaded file (CSV or XLSX)
    fieldnames, rows = _read_upload_rows(filename, csv_bytes)
    if fieldnames is None:
        raise HTTPException(status_code=400, detail="CSV file is empty or has no header row")

    actual_columns = {c.strip().lower() for c in fieldnames}
    missing = sorted(_PO_REQUIRED_COLUMNS - actual_columns)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Upload failed: missing required column(s): {', '.join(missing)}",
        )

    if not rows:
        raise HTTPException(status_code=400, detail="CSV file contains no data rows")

    # 2. Validate customer belongs to this organisation
    cust_result = await db.execute(
        select(Customer).where(
            Customer.id == customer_id,
            Customer.organisation_id == organisation_id,
        )
    )
    if cust_result.scalar_one_or_none() is None:
        raise HTTPException(status_code=404, detail="Customer not found")

    # 3. Duplicate detection (pre-check; IntegrityError on commit handles the concurrent case)
    existing = await db.execute(
        select(OutwardPO).where(
            OutwardPO.organisation_id == organisation_id,
            OutwardPO.po_number == po_number,
        )
    )
    if existing.scalar_one_or_none() is not None:
        raise HTTPException(
            status_code=409,
            detail="An outward PO already exists for this PO/Invoice number.",
        )

    # 4. Build lines — reject CSVs that contain rows for a different po_number
    lines_data: list[dict] = []
    for row in rows:
        row_po = row.get("po_number", "").strip()
        if row_po != po_number:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"CSV contains rows for a different PO number ('{row_po}'). "
                    f"All rows must match the requested PO '{po_number}'."
                ),
            )
        try:
            ordered_qty = int(row["ordered_qty"])
        except (ValueError, KeyError):
            raise HTTPException(status_code=400, detail="ordered_qty must be a positive integer")
        if ordered_qty <= 0:
            raise HTTPException(status_code=400, detail="ordered_qty must be a positive integer")
        lines_data.append({
            "ean": row["ean"].strip(),
            "ordered_qty": ordered_qty,
            "description": (row.get("description") or "").strip() or None,
        })

    if not lines_data:
        raise HTTPException(
            status_code=400,
            detail=f"No rows found in CSV for po_number '{po_number}'",
        )

    lines_data, _ = _consolidate_lines(lines_data)

    # 5. Create OutwardPO
    po = OutwardPO(
        organisation_id=organisation_id,
        customer_id=customer_id,
        po_number=po_number,
        status=OutwardPoStatusEnum.open,
        uploaded_by=uploaded_by,
    )
    db.add(po)
    await db.flush()  # populates po.id

    # 6. Create OutwardPOLine rows
    for ld in lines_data:
        db.add(OutwardPOLine(
            outward_po_id=po.id,
            organisation_id=organisation_id,
            ean=ld["ean"],
            description=ld["description"],
            ordered_qty=ld["ordered_qty"],
            packed_qty=0,
        ))

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="outward_po_uploaded",
        resource_type="outward_po",
        resource_id=po.id,
        user_id=uploaded_by,
        organisation_id=organisation_id,
        after_data={
            "po_number": po_number,
            "customer_id": customer_id,
            "line_count": len(lines_data),
        },
        ip_address=ip_address,
    )

    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(
            status_code=409,
            detail="An outward PO already exists for this PO/Invoice number.",
        )
    await db.refresh(po)
    return po


# ── Box: create ───────────────────────────────────────────────────────────────

async def create_box(
    db: AsyncSession,
    *,
    customer_id: int,
    organisation_id: int,
    created_by: int,
    ip_address: str | None,
) -> OutwardBox:
    from app.models.customer import Customer

    # 1. Load customer to get code for Box ID generation
    cust_result = await db.execute(
        select(Customer).where(
            Customer.id == customer_id,
            Customer.organisation_id == organisation_id,
        )
    )
    customer = cust_result.scalar_one_or_none()
    if customer is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Customer not found")

    # 2. Generate unique Box ID via atomic counter
    n = await _next_counter(
        db,
        counter_type=CounterTypeEnum.outward_box,
        organisation_id=organisation_id,
        customer_code=customer.code,
        date_key="",
    )
    box_id = f"OB-{customer.code}-{n:06d}"

    # 3. Persist
    box = OutwardBox(
        box_id=box_id,
        organisation_id=organisation_id,
        customer_id=customer_id,
        status=OutwardBoxStatusEnum.open,
        created_by=created_by,
    )
    db.add(box)
    await db.flush()  # populates box.id before writing the audit log

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="outward_box_created",
        resource_type="outward_box",
        resource_id=box.id,
        user_id=created_by,
        organisation_id=organisation_id,
        after_data={"box_id": box_id, "customer_id": customer_id},
        ip_address=ip_address,
    )

    await db.commit()
    await db.refresh(box)

    # Reload with scans eager-loaded
    result = await db.execute(
        select(OutwardBox)
        .options(selectinload(OutwardBox.scans))
        .where(OutwardBox.id == box.id)
    )
    return result.scalar_one()


# ── Box: get ──────────────────────────────────────────────────────────────────

async def get_box(
    db: AsyncSession,
    *,
    box_id: str,
    organisation_id: int,
) -> OutwardBox | None:
    result = await db.execute(
        select(OutwardBox)
        .options(selectinload(OutwardBox.scans))
        .where(
            OutwardBox.box_id == box_id,
            OutwardBox.organisation_id == organisation_id,
        )
    )
    return result.scalar_one_or_none()


# ── Box: close ────────────────────────────────────────────────────────────────

async def close_box(
    db: AsyncSession,
    *,
    box_id: str,
    organisation_id: int,
    closed_by: int,
    ip_address: str | None,
) -> OutwardBox:
    result = await db.execute(
        select(OutwardBox)
        .options(selectinload(OutwardBox.scans))
        .where(
            OutwardBox.box_id == box_id,
            OutwardBox.organisation_id == organisation_id,
        )
    )
    box = result.scalar_one_or_none()
    if box is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Box not found")

    if box.status == OutwardBoxStatusEnum.closed:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Box must be open or in_use to close (current: {box.status.value})",
        )

    before_status = box.status.value
    box.status = OutwardBoxStatusEnum.closed
    box.closed_by = closed_by
    box.closed_at = datetime.now(UTC)

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="outward_box_closed",
        resource_type="outward_box",
        resource_id=box.id,
        user_id=closed_by,
        organisation_id=organisation_id,
        before_data={"status": before_status},
        after_data={"status": "closed"},
        ip_address=ip_address,
    )

    await db.commit()

    # Reload with scans eager-loaded
    reloaded = await db.execute(
        select(OutwardBox)
        .options(selectinload(OutwardBox.scans))
        .where(OutwardBox.id == box.id)
    )
    return reloaded.scalar_one()


# ── Scan: add ─────────────────────────────────────────────────────────────────

async def _reject_scan(
    db: AsyncSession,
    *,
    organisation_id: int,
    box: OutwardBox,
    ean: str,
    created_by: int,
    ip_address: str | None,
    reason: str,
) -> None:
    """Persist a rejected scan attempt and audit it (OUT-12/R-3), then raise the
    PRD-mandated 400 with the exact reason text. The HTTP contract callers see is
    unchanged — only the rejection is now recorded instead of vanishing."""
    rejected_scan = OutwardScan(
        organisation_id=organisation_id,
        outward_box_id=box.id,
        outward_po_line_id=None,
        ean=ean,
        scan_result=OutwardScanResultEnum.rejected,
        reject_reason=reason,
        created_by=created_by,
    )
    db.add(rejected_scan)
    await db.flush()

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="outward_scan_rejected",
        resource_type="outward_scan",
        resource_id=rejected_scan.id,
        user_id=created_by,
        organisation_id=organisation_id,
        after_data={"box_id": box.box_id, "ean": ean, "reject_reason": reason},
        ip_address=ip_address,
    )

    await db.commit()
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=reason)


async def add_scan(
    db: AsyncSession,
    *,
    box_id: str,
    organisation_id: int,
    ean: str,
    created_by: int,
    ip_address: str | None,
) -> tuple[OutwardScan, str | None]:
    """Allocate EAN to oldest open outward PO line (FIFO) and insert a scan row.

    Returns (scan, note). note is the PRD 'no recorded inward stock' string or None.
    """
    # 1. Load box
    box_result = await db.execute(
        select(OutwardBox).where(
            OutwardBox.box_id == box_id,
            OutwardBox.organisation_id == organisation_id,
        )
    )
    box = box_result.scalar_one_or_none()
    if box is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Box not found")
    if box.status == OutwardBoxStatusEnum.closed:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cannot add scans to a closed box",
        )

    # 1b. One active box per packer (OUT-10). Advisory lock serialises concurrent
    # scans by the same packer into different boxes; the DB also carries a partial
    # unique index on (created_by) WHERE status='in_use' as the final backstop —
    # without this check that constraint surfaces as a raw 500 instead of this
    # PRD-mandated 400.
    if box.status == OutwardBoxStatusEnum.open:
        await db.execute(text("SELECT pg_advisory_xact_lock(:uid)"), {"uid": created_by})
        other_active = await db.execute(
            select(OutwardBox.id).where(
                OutwardBox.organisation_id == organisation_id,
                OutwardBox.created_by == created_by,
                OutwardBox.status == OutwardBoxStatusEnum.in_use,
                OutwardBox.id != box.id,
            )
        )
        if other_active.scalar_one_or_none() is not None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Mark the current box full before starting another box.",
            )

    # 2. Check if ANY open outward PO line has this EAN for this org/customer
    any_line_result = await db.execute(
        select(OutwardPOLine)
        .join(OutwardPO, OutwardPOLine.outward_po_id == OutwardPO.id)
        .where(
            OutwardPO.organisation_id == organisation_id,
            OutwardPO.customer_id == box.customer_id,
            OutwardPO.status == OutwardPoStatusEnum.open,
            OutwardPOLine.ean == ean,
        )
        .limit(1)
    )
    if any_line_result.scalar_one_or_none() is None:
        await _reject_scan(
            db,
            organisation_id=organisation_id,
            box=box,
            ean=ean,
            created_by=created_by,
            ip_address=ip_address,
            reason="EAN not found in open outward POs",
        )

    # 3. FIFO allocation with conditional UPDATE retry loop
    allocated_line_id: int | None = None
    for _ in range(_MAX_ALLOC_RETRIES):
        # Find oldest PO line with remaining capacity
        candidate_result = await db.execute(
            select(OutwardPOLine)
            .join(OutwardPO, OutwardPOLine.outward_po_id == OutwardPO.id)
            .where(
                OutwardPO.organisation_id == organisation_id,
                OutwardPO.customer_id == box.customer_id,
                OutwardPO.status == OutwardPoStatusEnum.open,
                OutwardPOLine.ean == ean,
                OutwardPOLine.packed_qty < OutwardPOLine.ordered_qty,
            )
            .order_by(OutwardPO.uploaded_at.asc(), OutwardPO.id.asc())
            .limit(1)
        )
        candidate = candidate_result.scalar_one_or_none()
        if candidate is None:
            await _reject_scan(
                db,
                organisation_id=organisation_id,
                box=box,
                ean=ean,
                created_by=created_by,
                ip_address=ip_address,
                reason="Quantity complete for all open outward POs",
            )

        # Conditional UPDATE — atomically increments only if still has capacity
        updated = await db.execute(
            text("""
                UPDATE outward_po_lines
                SET packed_qty = packed_qty + 1
                WHERE id = :id AND packed_qty < ordered_qty
                RETURNING id
            """),
            {"id": candidate.id},
        )
        if updated.scalar_one_or_none() is not None:
            allocated_line_id = candidate.id
            break
        # 0 rows updated — another session raced us; retry

    if allocated_line_id is None:
        await _reject_scan(
            db,
            organisation_id=organisation_id,
            box=box,
            ean=ean,
            created_by=created_by,
            ip_address=ip_address,
            reason="Quantity complete for all open outward POs",
        )

    # 4. Insert scan row
    scan = OutwardScan(
        organisation_id=organisation_id,
        outward_box_id=box.id,
        outward_po_line_id=allocated_line_id,
        ean=ean,
        scan_result=OutwardScanResultEnum.accepted,
        created_by=created_by,
    )
    db.add(scan)

    # 5. If box.status == open, transition to in_use
    if box.status == OutwardBoxStatusEnum.open:
        await db.execute(
            text("UPDATE outward_boxes SET status = 'in_use' WHERE id = :id"),
            {"id": box.id},
        )

    # 6. Flush scan to get its id before the ledger entry
    await db.flush()

    # 7. Write -1 InventoryLedgerEntry (outgoing stock)
    db.add(InventoryLedgerEntry(
        organisation_id=organisation_id,
        customer_id=box.customer_id,
        ean=ean,
        quantity_change=-1,
        source_type=LedgerSourceTypeEnum.outward_scan,
        source_id=scan.id,
    ))
    await db.flush()

    # 8. Check ledger balance for "no inward stock" note
    ledger_result = await db.execute(
        text("""
            SELECT COALESCE(SUM(quantity_change), 0)
            FROM inventory_ledger_entries
            WHERE organisation_id = :org AND customer_id = :cust AND ean = :ean
        """),
        {"org": organisation_id, "cust": box.customer_id, "ean": ean},
    )
    note: str | None = None
    if (ledger_result.scalar() or 0) <= 0:
        note = "Note: no recorded inward stock for this item."
    scan.stock_flagged = note is not None

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="outward_scan_added",
        resource_type="outward_scan",
        resource_id=scan.id,
        user_id=created_by,
        organisation_id=organisation_id,
        after_data={"box_id": box_id, "ean": ean, "po_line_id": allocated_line_id},
        ip_address=ip_address,
    )

    await db.commit()
    await db.refresh(scan)
    return scan, note


# ── Scan: delete ──────────────────────────────────────────────────────────────

async def delete_scan(
    db: AsyncSession,
    *,
    scan_id: int,
    organisation_id: int,
    deleted_by: int,
    ip_address: str | None,
) -> OutwardScan:
    """Soft-delete an outward scan. Decrements packed_qty and writes +1 ledger reversal."""
    scan_result = await db.execute(
        select(OutwardScan)
        .join(OutwardBox, OutwardScan.outward_box_id == OutwardBox.id)
        .where(
            OutwardScan.id == scan_id,
            OutwardBox.organisation_id == organisation_id,
        )
    )
    scan = scan_result.scalar_one_or_none()
    if scan is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Scan not found")
    if scan.scan_result == OutwardScanResultEnum.deleted:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="Scan already deleted"
        )
    if scan.scan_result == OutwardScanResultEnum.rejected:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="Cannot delete a rejected scan"
        )

    # Load box for customer_id (needed for ledger entry)
    box_result = await db.execute(
        select(OutwardBox).where(
            OutwardBox.id == scan.outward_box_id,
            OutwardBox.organisation_id == organisation_id,
        )
    )
    box = box_result.scalar_one()

    if box.status == OutwardBoxStatusEnum.closed:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cannot delete scans from a closed box",
        )
    if box.created_by != deleted_by:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You can only delete scans from your own active box.",
        )

    # Soft-delete
    scan.scan_result = OutwardScanResultEnum.deleted
    scan.deleted_at = datetime.now(UTC)
    scan.deleted_by = deleted_by

    # Decrement po_line packed_qty (if allocated)
    if scan.outward_po_line_id is not None:
        await db.execute(
            text("UPDATE outward_po_lines SET packed_qty = packed_qty - 1 WHERE id = :id"),
            {"id": scan.outward_po_line_id},
        )

    # Write +1 ledger reversal
    db.add(InventoryLedgerEntry(
        organisation_id=organisation_id,
        customer_id=box.customer_id,
        ean=scan.ean,
        quantity_change=1,
        source_type=LedgerSourceTypeEnum.outward_scan_deletion,
        source_id=scan.id,
    ))

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="outward_scan_deleted",
        resource_type="outward_scan",
        resource_id=scan.id,
        user_id=deleted_by,
        organisation_id=organisation_id,
        before_data={"ean": scan.ean, "scan_result": "accepted"},
        after_data={"scan_result": "deleted"},
        ip_address=ip_address,
    )

    await db.commit()
    await db.refresh(scan)
    return scan


# ── PO: list open ─────────────────────────────────────────────────────────────

async def list_open_pos(
    db: AsyncSession,
    *,
    org_id: int,
    customer_id: int | None = None,
    search: str | None = None,
    from_date: date | None = None,
    to_date: date | None = None,
    status_filter: OutwardPoStatusEnum | None = None,
    page: int = 1,
    page_size: int = 20,
) -> OpenPOListResponse:
    from app.models.customer import Customer

    base_where = [OutwardPO.organisation_id == org_id]
    if status_filter is not None:
        base_where.append(OutwardPO.status == status_filter)
    if customer_id is not None:
        base_where.append(OutwardPO.customer_id == customer_id)
    if search:
        base_where.append(OutwardPO.po_number.ilike(f"%{search}%"))
    if from_date is not None:
        start = datetime(from_date.year, from_date.month, from_date.day, tzinfo=UTC)
        base_where.append(OutwardPO.uploaded_at >= start)
    if to_date is not None:
        end = datetime(to_date.year, to_date.month, to_date.day, tzinfo=UTC) + timedelta(days=1)
        base_where.append(OutwardPO.uploaded_at < end)

    total = await db.scalar(
        select(sql_func.count()).select_from(OutwardPO).where(*base_where)
    ) or 0

    result = await db.execute(
        select(
            OutwardPO,
            Customer.name.label("customer_name"),
            sql_func.coalesce(sql_func.sum(OutwardPOLine.ordered_qty), 0).label("total_ordered"),
            sql_func.coalesce(sql_func.sum(OutwardPOLine.packed_qty), 0).label("total_packed"),
        )
        .join(Customer, OutwardPO.customer_id == Customer.id)
        .outerjoin(OutwardPOLine, OutwardPOLine.outward_po_id == OutwardPO.id)
        .where(*base_where)
        .group_by(OutwardPO.id, Customer.name)
        .order_by(OutwardPO.uploaded_at.desc())
        .limit(page_size)
        .offset((page - 1) * page_size)
    )
    rows = result.all()

    items = []
    for po, customer_name, total_ordered, total_packed in rows:
        progress_pct = (total_packed / total_ordered * 100.0) if total_ordered > 0 else 0.0
        items.append(OpenPOSummary(
            id=po.id,
            po_number=po.po_number,
            customer_id=po.customer_id,
            customer_name=customer_name,
            status=po.status,
            uploaded_at=po.uploaded_at,
            total_ordered=total_ordered,
            total_packed=total_packed,
            progress_pct=round(progress_pct, 1),
        ))
    return OpenPOListResponse(items=items, total=total)


# ── PO: toggle status ─────────────────────────────────────────────────────────

async def toggle_po_status(
    db: AsyncSession,
    *,
    po_id: int,
    org_id: int,
    new_status: OutwardPoStatusEnum,
    user_id: int,
    ip_address: str | None,
) -> OutwardPO:
    result = await db.execute(
        select(OutwardPO).where(
            OutwardPO.id == po_id,
            OutwardPO.organisation_id == org_id,
        )
    )
    po = result.scalar_one_or_none()
    if po is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="PO not found")
    if po.status == new_status:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"PO is already {new_status.value}.",
        )

    old_status = po.status
    po.status = new_status

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="po_status_changed",
        resource_type="outward_po",
        resource_id=po.id,
        user_id=user_id,
        organisation_id=org_id,
        before_data={"status": old_status.value},
        after_data={"status": new_status.value},
        ip_address=ip_address,
    )

    await db.commit()

    result = await db.execute(
        select(OutwardPO).options(selectinload(OutwardPO.lines)).where(OutwardPO.id == po.id)
    )
    return result.scalar_one()


# ── PO: preview (stateless parse) ────────────────────────────────────────────

async def preview_po(
    db: AsyncSession,
    *,
    org_id: int,
    customer_id: int,
    filename: str = "po.csv",
    csv_bytes: bytes,
) -> OutwardPOPreviewResponse:
    from app.models.customer import Customer

    problems: list[str] = []

    # Validate customer belongs to org
    cust_result = await db.execute(
        select(Customer).where(
            Customer.id == customer_id,
            Customer.organisation_id == org_id,
        )
    )
    if cust_result.scalar_one_or_none() is None:
        problems.append("Customer not found in this organisation.")
        return OutwardPOPreviewResponse(
            first_10_rows=[], total_rows=0, problems=problems, is_valid=False
        )

    # Parse the uploaded file (CSV or XLSX)
    try:
        fieldnames, rows = _read_upload_rows(filename, csv_bytes)
    except Exception:
        problems.append("File encoding error — upload a UTF-8 CSV or a valid XLSX file.")
        return OutwardPOPreviewResponse(
            first_10_rows=[], total_rows=0, problems=problems, is_valid=False
        )

    if fieldnames is None:
        problems.append("CSV file is empty or has no header row.")
        return OutwardPOPreviewResponse(
            first_10_rows=[], total_rows=0, problems=problems, is_valid=False
        )

    actual_columns = {c.strip().lower() for c in fieldnames}
    missing = sorted(_PO_REQUIRED_COLUMNS - actual_columns)
    if missing:
        problems.append(f"Upload failed: missing required column(s): {', '.join(missing)}")
        return OutwardPOPreviewResponse(
            first_10_rows=[], total_rows=0, problems=problems, is_valid=False
        )

    if not rows:
        problems.append("CSV file contains no data rows.")
        return OutwardPOPreviewResponse(
            first_10_rows=[], total_rows=0, problems=problems, is_valid=False
        )

    po_numbers = {row.get("po_number", "").strip() for row in rows}
    if len(po_numbers) > 1:
        problems.append(
            f"CSV contains {len(po_numbers)} different PO numbers. "
            "Each upload must contain rows for a single PO."
        )

    preview_rows: list[OutwardPOPreviewRow] = []
    for i, row in enumerate(rows[:10]):
        try:
            qty = int(row.get("ordered_qty", ""))
        except (ValueError, TypeError):
            problems.append(f"Row {i + 1}: ordered_qty is not a valid integer.")
            qty = 0
        preview_rows.append(OutwardPOPreviewRow(
            po_number=row.get("po_number", "").strip(),
            ean=row.get("ean", "").strip(),
            ordered_qty=qty,
            description=(row.get("description") or "").strip() or None,
        ))

    # OUT-2: total quantity across the whole file (not just the first-10 preview)
    total_quantity = 0
    for i, row in enumerate(rows):
        try:
            total_quantity += int(row.get("ordered_qty", ""))
        except (ValueError, TypeError):
            if i >= 10:  # rows 0-9 are already flagged by the preview_rows loop above
                problems.append(f"Row {i + 1}: ordered_qty is not a valid integer.")

    # OUT-3/OUT-2: warn (never block) when duplicate EANs will be consolidated
    ean_counts: dict[str, int] = {}
    for row in rows:
        ean = row.get("ean", "").strip()
        if ean:
            ean_counts[ean] = ean_counts.get(ean, 0) + 1
    duplicate_eans = {ean: n for ean, n in ean_counts.items() if n > 1}
    consolidation_notice = (
        f"{sum(duplicate_eans.values())} row(s) across {len(duplicate_eans)} duplicate EAN(s) "
        "will be consolidated into a single line each, with quantities summed."
        if duplicate_eans
        else None
    )

    return OutwardPOPreviewResponse(
        first_10_rows=preview_rows,
        total_rows=len(rows),
        total_quantity=total_quantity,
        problems=problems,
        is_valid=len(problems) == 0,
        consolidation_notice=consolidation_notice,
    )


# ── Labels: generate ──────────────────────────────────────────────────────────

async def generate_labels(
    db: AsyncSession,
    *,
    org_id: int,
    customer_id: int,
    count: int,
    created_by: int,
    ip_address: str | None,
) -> LabelGenerateResponse:
    from app.models.customer import Customer

    cust_result = await db.execute(
        select(Customer).where(
            Customer.id == customer_id,
            Customer.organisation_id == org_id,
        )
    )
    customer = cust_result.scalar_one_or_none()
    if customer is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Customer not found")

    box_ids: list[str] = []
    for _ in range(count):
        n = await _next_counter(
            db,
            counter_type=CounterTypeEnum.outward_box,
            organisation_id=org_id,
            customer_code=customer.code,
            date_key="",
        )
        box_id = f"OB-{customer.code}-{n:06d}"
        box = OutwardBox(
            box_id=box_id,
            organisation_id=org_id,
            customer_id=customer_id,
            status=OutwardBoxStatusEnum.open,
            print_count=1,
            created_by=created_by,
        )
        db.add(box)
        await db.flush()
        box_ids.append(box_id)

    await write_audit_log(
        db,
        module=AuditModuleEnum.outward,
        action="labels_generated",
        resource_type="outward_box",
        resource_id=None,
        user_id=created_by,
        organisation_id=org_id,
        after_data={"customer_id": customer_id, "count": count, "box_ids": box_ids},
        ip_address=ip_address,
    )

    await db.commit()
    return LabelGenerateResponse(box_ids=box_ids)


# ── Labels: history ───────────────────────────────────────────────────────────

async def get_label_history(
    db: AsyncSession,
    *,
    org_id: int,
    user_id: int,
    customer_id: int | None = None,
    page: int = 1,
    page_size: int = 20,
) -> LabelHistoryResponse:
    from app.models.customer import Customer

    base_where = [
        OutwardBox.organisation_id == org_id,
        OutwardBox.created_by == user_id,
    ]
    if customer_id is not None:
        base_where.append(OutwardBox.customer_id == customer_id)

    total = await db.scalar(
        select(sql_func.count()).select_from(OutwardBox).where(*base_where)
    ) or 0

    result = await db.execute(
        select(OutwardBox, Customer.name.label("customer_name"))
        .join(Customer, OutwardBox.customer_id == Customer.id)
        .where(*base_where)
        .order_by(OutwardBox.created_at.desc())
        .limit(page_size)
        .offset((page - 1) * page_size)
    )
    rows = result.all()

    return LabelHistoryResponse(
        items=[
            LabelHistoryItem(
                box_id=box.box_id,
                customer_name=customer_name,
                status=box.status,
                print_count=box.print_count,
                created_at=box.created_at,
                closed_at=box.closed_at,
            )
            for box, customer_name in rows
        ],
        total=total,
    )


# ── Boxes: packing history ────────────────────────────────────────────────────

async def list_packing_history(
    db: AsyncSession,
    *,
    org_id: int,
    user_id: int,
    days: int = 30,
    page: int = 1,
    page_size: int = 20,
) -> OutwardBoxListResponse:
    from app.models.customer import Customer

    since = datetime.now(UTC) - timedelta(days=days)
    base_where = [
        OutwardBox.organisation_id == org_id,
        OutwardBox.created_by == user_id,
        OutwardBox.created_at >= since,
    ]

    total = await db.scalar(
        select(sql_func.count()).select_from(OutwardBox).where(*base_where)
    ) or 0

    # Subquery: count accepted scans per box
    scan_count_sq = (
        select(sql_func.count())
        .select_from(OutwardScan)
        .where(
            OutwardScan.outward_box_id == OutwardBox.id,
            OutwardScan.scan_result == OutwardScanResultEnum.accepted,
        )
        .correlate(OutwardBox)
        .scalar_subquery()
    )

    result = await db.execute(
        select(OutwardBox, Customer.name.label("customer_name"), scan_count_sq.label("scan_count"))
        .join(Customer, OutwardBox.customer_id == Customer.id)
        .where(*base_where)
        .order_by(OutwardBox.created_at.desc())
        .limit(page_size)
        .offset((page - 1) * page_size)
    )
    rows = result.all()

    return OutwardBoxListResponse(
        items=[
            OutwardBoxSummary(
                box_id=box.box_id,
                customer_name=customer_name,
                status=box.status,
                scan_count=scan_count,
                closed_at=box.closed_at,
                created_at=box.created_at,
            )
            for box, customer_name, scan_count in rows
        ],
        total=total,
    )
