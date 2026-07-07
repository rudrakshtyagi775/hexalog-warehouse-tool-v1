from datetime import date, timedelta

import pytest_asyncio
from httpx import AsyncClient

from app.models.customer import Customer
from app.models.enums import CustomerStatusEnum


async def _login(client: AsyncClient, email: str, password: str, org_id: int) -> str:
    resp = await client.post(
        "/api/auth/login",
        json={"email": email, "password": password, "organisation_id": org_id},
    )
    return resp.json()["access_token"]


def _make_csv(po_number: str, ean: str = "1234567890123", qty: int = 5) -> bytes:
    return f"po_number,ean,ordered_qty,description\n{po_number},{ean},{qty},Widget\n".encode()


@pytest_asyncio.fixture
async def customer(db, org) -> Customer:
    c = Customer(
        organisation_id=org.id, name="Test Customer", code="TST",
        status=CustomerStatusEnum.active,
    )
    db.add(c)
    await db.flush()
    return c


async def _submit_one_box(client, inward_token, org, customer, reference_id, box_number, ean):
    box_resp = await client.post(
        "/api/inward/boxes",
        json={
            "customer_id": customer.id,
            "inward_reference_id": reference_id,
            "box_number": box_number,
        },
        headers={"Authorization": f"Bearer {inward_token}"},
    )
    box_id = box_resp.json()["box_id"]

    await client.post(
        f"/api/inward/boxes/{box_id}/scans",
        json={"ean": ean},
        headers={"Authorization": f"Bearer {inward_token}"},
    )
    await client.post(
        f"/api/inward/boxes/{box_id}/close",
        json={"physical_qty": 1},
        headers={"Authorization": f"Bearer {inward_token}"},
    )
    submit_resp = await client.post(
        f"/api/inward/boxes/{box_id}/submit",
        headers={"Authorization": f"Bearer {inward_token}"},
    )
    return submit_resp.json()


async def test_inscan_report_returns_row_per_ean_per_box(
    client, admin_user, inward_operator_user, org, customer
):
    inward_token = await _login(client, "inward@test.com", "InwardPass1!", org.id)
    ean = "1234567890123"
    await client.post(
        "/api/inward/pos",
        data={"po_number": "IN-REPORT-001", "customer_id": customer.id},
        files={"file": ("po.csv", _make_csv("IN-REPORT-001", ean=ean), "text/csv")},
        headers={"Authorization": f"Bearer {inward_token}"},
    )
    reference_id = (await client.post(
        "/api/inward/references",
        json={"customer_id": customer.id, "po_number": "IN-REPORT-001"},
        headers={"Authorization": f"Bearer {inward_token}"},
    )).json()["id"]

    submitted = await _submit_one_box(
        client, inward_token, org, customer, reference_id, "1", ean
    )

    admin_token = await _login(client, "admin@test.com", "AdminPass1!", org.id)
    today = date.today()
    resp = await client.get(
        "/api/admin/reports/inscan",
        params={"from_date": str(today - timedelta(days=1)), "to_date": str(today)},
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    row = body["items"][0]
    assert row["inscan_number"] == submitted["inscan_number"]
    assert row["customer_name"] == "Test Customer"
    assert row["po_number"] == "IN-REPORT-001"
    assert row["box_number"] == "1"
    assert row["ean"] == ean
    assert row["scanned_qty"] == 1
    assert row["physical_qty"] == 1
    assert row["variance"] == 0
    assert row["status"] == "completed"


async def test_inscan_report_csv_format(client, admin_user, inward_operator_user, org, customer):
    inward_token = await _login(client, "inward@test.com", "InwardPass1!", org.id)
    ean = "9999999999999"
    await client.post(
        "/api/inward/pos",
        data={"po_number": "IN-REPORT-CSV", "customer_id": customer.id},
        files={"file": ("po.csv", _make_csv("IN-REPORT-CSV", ean=ean), "text/csv")},
        headers={"Authorization": f"Bearer {inward_token}"},
    )
    reference_id = (await client.post(
        "/api/inward/references",
        json={"customer_id": customer.id, "po_number": "IN-REPORT-CSV"},
        headers={"Authorization": f"Bearer {inward_token}"},
    )).json()["id"]
    await _submit_one_box(client, inward_token, org, customer, reference_id, "1", ean)

    admin_token = await _login(client, "admin@test.com", "AdminPass1!", org.id)
    today = date.today()
    resp = await client.get(
        "/api/admin/reports/inscan",
        params={
            "from_date": str(today - timedelta(days=1)),
            "to_date": str(today),
            "format": "csv",
        },
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "Inscan Number" in resp.text
    assert ean in resp.text


async def test_inscan_report_xlsx_format(client, admin_user, inward_operator_user, org, customer):
    """Reports intro: 'CSV and XLSX download' — verify XLSX round-trips real data."""
    from io import BytesIO

    from openpyxl import load_workbook

    inward_token = await _login(client, "inward@test.com", "InwardPass1!", org.id)
    ean = "5555555555555"
    await client.post(
        "/api/inward/pos",
        data={"po_number": "IN-REPORT-XLSX", "customer_id": customer.id},
        files={"file": ("po.csv", _make_csv("IN-REPORT-XLSX", ean=ean), "text/csv")},
        headers={"Authorization": f"Bearer {inward_token}"},
    )
    reference_id = (await client.post(
        "/api/inward/references",
        json={"customer_id": customer.id, "po_number": "IN-REPORT-XLSX"},
        headers={"Authorization": f"Bearer {inward_token}"},
    )).json()["id"]
    await _submit_one_box(client, inward_token, org, customer, reference_id, "1", ean)

    admin_token = await _login(client, "admin@test.com", "AdminPass1!", org.id)
    today = date.today()
    resp = await client.get(
        "/api/admin/reports/inscan",
        params={
            "from_date": str(today - timedelta(days=1)),
            "to_date": str(today),
            "format": "xlsx",
        },
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "Inscan Number" in header
    data_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert ean in data_row


async def test_inscan_report_requires_admin(client, inward_operator_user, org):
    token = await _login(client, "inward@test.com", "InwardPass1!", org.id)
    today = date.today()
    resp = await client.get(
        "/api/admin/reports/inscan",
        params={"from_date": str(today), "to_date": str(today)},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 403


async def test_inscan_report_empty_when_no_completed_boxes(client, admin_user, org):
    token = await _login(client, "admin@test.com", "AdminPass1!", org.id)
    today = date.today()
    resp = await client.get(
        "/api/admin/reports/inscan",
        params={"from_date": str(today), "to_date": str(today)},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 200
    assert resp.json()["total"] == 0


async def test_inscan_report_variance_sign_is_physical_minus_scanned(
    client, admin_user, org, customer, db
):
    """PRD 9.1: Variance = physical - scanned, not scanned - physical.

    Constructs a completed box directly (bypassing the app's submit-time
    scanned==physical equality rule) purely to pin down the report's arithmetic.
    """
    from datetime import UTC, datetime

    from app.models.enums import InwardBoxStatusEnum
    from app.models.inward import InwardBox, InwardScan

    box = InwardBox(
        box_id="B-TST-000099",
        organisation_id=org.id,
        customer_id=customer.id,
        status=InwardBoxStatusEnum.completed,
        scanned_qty=5,
        physical_qty=3,
        inscan_number="INS-TST-99999999-0001",
        submitted_at=datetime.now(UTC),
    )
    db.add(box)
    await db.flush()
    db.add(InwardScan(
        organisation_id=org.id, inward_box_id=box.id, ean="9990009900099",
        is_deleted=False,
    ))
    await db.commit()

    token = await _login(client, "admin@test.com", "AdminPass1!", org.id)
    today = date.today()
    resp = await client.get(
        "/api/admin/reports/inscan",
        params={"from_date": str(today), "to_date": str(today)},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 200
    row = resp.json()["items"][0]
    assert row["physical_qty"] == 3
    assert row["scanned_qty"] == 1  # per-EAN count, not the box-level scanned_qty
    assert row["variance"] == 3 - 5  # physical(3) - box-level scanned(5) == -2
